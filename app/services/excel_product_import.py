from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from app.services.error_logging import get_error_logger, log_failures

from app.config import CATEGORIES, CATEGORIES_AR


@dataclass
class ProductImportRow:
    row_number: int
    code: str
    name: str
    category: str
    supplier: str
    ah: float | None = None
    l: float | None = None
    w: float | None = None
    h: float | None = None
    oh: float | None = None
    layout: int | None = None
    pattern: str = ""
    pallet: int | None = None


@dataclass
class ProductImportReport:
    rows: list[ProductImportRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


class ExcelProductImporter:
    COLUMNS = {
        "code": ["product", "code", "product code", "code / name", "الكود / الاسم", "كود", "كود الصنف", "رمز الصنف"],
        "name": ["product", "name", "product name", "tire name", "اسم الإطار", "اسم الصنف", "اسم المنتج", "الصنف"],
        "category": ["category", "التصنيف", "الفئة"],
        "supplier": ["supplier", "supplier name", "supplier code", "المورد", "اسم المورد", "كود المورد"],
        "pattern": ["pattern", "tread", "tread pattern", "النقشة", "نقشة"],
        "ah": ["ah"],
        "l": ["l", "length", "الطول"],
        "w": ["w", "width", "العرض"],
        "h": ["h", "height", "الارتفاع"],
        "oh": ["oh"],
        "layout": ["layout", "layout code", "التخطيط", "رمز التخطيط"],
        "pallet": ["pallet", "الطبلية"],
    }
    REQUIRED = ("category", "supplier")
    NUMERIC = ("ah", "l", "w", "h", "oh", "layout", "pallet")

    @staticmethod
    def _normalize(value: Any) -> str:
        return "" if value is None else str(value).strip()

    @classmethod
    def generate_template(cls, file_path: str) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Products"
        sheet.append([
            "code", "category", "supplier",
            "ah", "l", "w", "h", "oh", "layout", "pallet", "name", "pattern",
        ])
        sheet.append([
            "YOUR-CODE", CATEGORIES[0],
            "Existing Supplier Code or Name", 100, 200, 50, 60, 0, 1, 20, "", "",
        ])
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = max(
                14, min(30, max(len(str(cell.value or "")) for cell in column) + 2)
            )
        workbook.save(file_path)

    @classmethod
    def import_file(cls, file_path: str) -> ProductImportReport:
        report = ProductImportReport()
        try:
            workbook = load_workbook(file_path, data_only=True, read_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                report.errors.append("ملف Excel فارغ أو لا يحتوي على رؤوس أعمدة.")
                workbook.close()
                return report

            normalized_headers = {
                cls._normalize(value).casefold(): index
                for index, value in enumerate(headers)
                if cls._normalize(value)
            }
            mapped = {}
            for key, aliases in cls.COLUMNS.items():
                mapped[key] = next(
                    (normalized_headers[alias.casefold()] for alias in aliases
                     if alias.casefold() in normalized_headers),
                    None,
                )

            missing = [key for key in cls.REQUIRED if mapped[key] is None]
            if mapped["code"] is None and mapped["name"] is None:
                missing.append("code / name")
            if missing:
                report.errors.append("أعمدة مطلوبة مفقودة: " + ", ".join(missing))
                workbook.close()
                return report

            category_map = {
                value.casefold(): CATEGORIES[index]
                for index, value in enumerate(CATEGORIES)
            }
            category_map.update({
                value.casefold(): CATEGORIES[index]
                for index, value in enumerate(CATEGORIES_AR)
            })
            for row_number, values in enumerate(rows, start=2):
                if not values or all(not cls._normalize(value) for value in values):
                    continue
                data = {
                    key: values[index] if index is not None and index < len(values) else None
                    for key, index in mapped.items()
                }
                code = cls._normalize(data["code"])
                name = cls._normalize(data["name"]) or code
                supplier = cls._normalize(data["supplier"])
                category_raw = cls._normalize(data["category"])
                category = category_map.get(category_raw.casefold())
                pattern = cls._normalize(data["pattern"])
                identity = code if category == "Batteries" else name
                if not identity or not supplier:
                    report.errors.append(f"السطر {row_number}: كود البطارية أو اسم الإطار والمورد مطلوبة.")
                    continue
                if not category:
                    report.errors.append(
                        f"السطر {row_number}: تصنيف غير معروف ({category_raw})."
                    )
                    continue

                if category != "Batteries" and not pattern:
                    report.errors.append(f"السطر {row_number}: نقشة الإطار مطلوبة.")
                    continue
                code = identity
                name = identity
                numeric_values = {}
                invalid = False
                for key in cls.NUMERIC:
                    value = data[key] if category == "Batteries" else None
                    if value is None or cls._normalize(value) == "":
                        numeric_values[key] = None
                        continue
                    try:
                        number = float(str(value).replace(",", ""))
                        if key in ("layout", "pallet"):
                            if number < 0 or not number.is_integer():
                                raise ValueError
                            number = int(number)
                        numeric_values[key] = number
                    except (TypeError, ValueError):
                        report.errors.append(
                            f"السطر {row_number}: قيمة غير رقمية أو غير صالحة في {key}."
                        )
                        invalid = True
                        break
                if invalid:
                    continue

                report.rows.append(ProductImportRow(
                    row_number=row_number,
                    code=code,
                    name=name,
                    category=category,
                    supplier=supplier,
                    pattern=pattern,
                    **numeric_values,
                ))
            workbook.close()
        except Exception as exc:
            get_error_logger("product_import").exception("Failed to read product spreadsheet")
            report.errors.append(f"تعذر قراءة ملف Excel: {exc}")
        return report

    @classmethod
    @log_failures("save imported products")
    def save_to_database(cls, report: ProductImportReport) -> tuple[int, list[str]]:
        from app.services.product_service import product_service
        from app.services.supplier_service import supplier_service

        suppliers = supplier_service.get_all()
        supplier_map = {}
        for supplier in suppliers:
            supplier_map[supplier["name"].strip().casefold()] = supplier["id"]
            supplier_map[supplier["code"].strip().casefold()] = supplier["id"]

        created = 0
        errors = list(report.errors)
        for row in report.rows:
            supplier_id = supplier_map.get(row.supplier.casefold())
            if supplier_id is None:
                errors.append(
                    f"السطر {row.row_number}: المورد '{row.supplier}' غير موجود؛ "
                    "استخدم اسم أو كود مورد مسجل."
                )
                continue
            try:
                product_id = product_service.create(
                    code=row.code,
                    name=row.name,
                    pattern=row.pattern,
                    category=row.category,
                    supplier_id=supplier_id,
                    ah=row.ah,
                    l=row.l,
                    w=row.w,
                    h=row.h,
                    oh=row.oh,
                    layout=row.layout,
                    pallet=row.pallet,
                )
                if not product_service.get_by_id(product_id):
                    errors.append(
                        f"السطر {row.row_number} ({row.code}): تعذر تأكيد حفظ الصنف."
                    )
                    continue
                created += 1
            except Exception as exc:
                get_error_logger("product_import").warning(
                    "Skipped product import row %s: %s", row.row_number, exc,
                )
                errors.append(f"السطر {row.row_number} ({row.code}): {exc}")
        return created, errors
