from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from app.services.error_logging import get_error_logger, log_failures


@dataclass
class InvoiceRow:
    invoice_no: str
    date: str
    supplier: str
    product: str
    qty: float
    unit_price: float
    purchase_price: float = 0.0
    sale_price: float = 0.0
    tax: float = 0.0
    discount: float = 0.0
    notes: str = ""
    product_code: str = ""
    category: str = ""
    pattern: str = ""
    pallet: int | None = None
    row_number: int = 0


@dataclass
class ImportReport:
    rows: list[InvoiceRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


class ExcelInvoiceImporter:
    REQUIRED_COLUMNS = {
        "invoice_no": ["invoice_no", "invoice no", "رقم الفاتورة", "رقم_الفاتورة", "الفاتورة", "bill_no"],
        "date": ["date", "التاريخ", "date_invoice", "تاريخ"],
        "supplier": ["supplier", "supplier_name", "المورد", "اسم المورد", "vendor"],
        "product": ["product", "product_name", "name", "tire name", "اسم الإطار", "اسم الصنف", "الصنف", "المنتج", "اسم المنتج", "item"],
        "product_code": ["product_code", "product code", "item_code", "code", "كود الصنف", "رمز الصنف"],
        "category": ["category", "التصنيف"],
        "pattern": ["pattern", "tread", "النقشة", "نقشة"],
        "qty": ["qty", "quantity", "الكمية", "qty_ordered", "العدد"],
        "purchase_price": ["purchase_price", "purchase price", "السعر_الشراء", "سعر الشراء", "purchase_usd", "buy_price", "price_purchase", "purchase"],
        "sale_price": ["sale_price", "sale price", "السعر_البيع", "سعر البيع", "selling_price", "sell_price", "sale", "selling price", "unit_price", "price", "السعر", "unit price", "سعر الوحدة"],
        "pallet": ["pallet", "pallets", "الطبلية", "طبلية", "عدد الطبالي", "عدد الطبلية"],
    }

    @staticmethod
    def generate_template(file_path: str) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Invoices"
        headers = [
            "invoice_no",
            "date",
            "supplier",
            "product",
            "pallet",
            "qty",
            "unit_price",
            "category",
        ]
        ws.append(headers)
        ws.append(["INV-2026-0001", date.today().isoformat(), "Supplier Example", "YOUR-CODE", 1, 10, 120, "Batteries"])
        ws.freeze_panes = "A2"
        wb.save(file_path)

    @staticmethod
    def normalize(value: Any) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def _parse_date(value) -> str:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        text = str(value or "").strip()
        if not text:
            return date.today().isoformat()
        try:
            return datetime.fromisoformat(text).date().isoformat()
        except ValueError:
            for pattern in ("%d/%m/%Y", "%d-%m-%Y"):
                try:
                    return datetime.strptime(text, pattern).date().isoformat()
                except ValueError:
                    pass
        raise ValueError("تاريخ الفاتورة غير صالح؛ استخدم YYYY-MM-DD.")

    @staticmethod
    def _positive(value) -> float:
        number = Decimal(str(value or "0").strip())
        if not number.is_finite() or number <= 0:
            raise ValueError("الكميات والأسعار يجب أن تكون أرقامًا موجبة.")
        result = float(number)
        import math
        if not math.isfinite(result) or result <= 0:
            raise ValueError("القيمة أكبر من المسموح.")
        return result

    @staticmethod
    def _parse_pallet(value: Any) -> int | None:
        if value is None or str(value).strip() == "":
            return None
        try:
            number = Decimal(str(value).strip())
            if not number.is_finite() or number < 0 or number != number.to_integral_value():
                raise ValueError
            return int(number)
        except (InvalidOperation, ValueError):
            raise ValueError("الطبلية يجب أن تكون عددًا صحيحًا غير سالب.") from None

    @classmethod
    def _find_column(cls, headers: Iterable[str], candidates: Iterable[str]) -> str | None:
        normalized = {h.strip().lower(): h for h in headers if h is not None}
        for candidate in candidates:
            key = candidate.strip().lower()
            if key in normalized:
                return normalized[key]
        return None

    @classmethod
    def import_file(cls, file_path: str) -> ImportReport:
        report = ImportReport()
        workbook = None

        try:
            workbook = load_workbook(file_path, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))

            if not rows:
                report.errors.append("الملف فارغ.")
                return report

            headers = [cls.normalize(v) for v in rows[0]]
            if not headers:
                report.errors.append("لا توجد رؤوس أعمدة في الملف.")
                return report

            mapped = {}
            for field_name, candidates in cls.REQUIRED_COLUMNS.items():
                mapped[field_name] = cls._find_column(headers, candidates)

            missing = [
                name for name, value in mapped.items()
                if value is None and name not in {"sale_price", "purchase_price", "product_code", "product", "category", "pattern", "pallet"}
            ]
            if mapped["product"] is None and mapped["product_code"] is None:
                missing.append("product_code")
            if mapped["purchase_price"] is None and mapped["sale_price"] is None:
                missing.append("purchase_price")
            if missing:
                report.errors.append("الأعمدة التالية مفقودة: " + ", ".join(missing))
                return report

            purchase_col = mapped["purchase_price"]
            sale_col = mapped["sale_price"]

            for row_number, row in enumerate(rows[1:], start=2):
                if row is None or all(cls.normalize(v) == "" for v in row):
                    continue

                values = {
                    header: row[idx] if idx < len(row) else ""
                    for idx, header in enumerate(headers)
                }

                try:
                    purchase_price = cls._positive(values[purchase_col]) if purchase_col else None
                    sale_price = cls._positive(values[sale_col]) if sale_col else None
                    if purchase_col is None:
                        purchase_price = sale_price
                    if sale_col is None:
                        sale_price = purchase_price

                    item = InvoiceRow(
                        invoice_no=cls.normalize(values.get(mapped["invoice_no"], "")),
                        date=cls._parse_date(values.get(mapped["date"], "")),
                        supplier=cls.normalize(values.get(mapped["supplier"], "")),
                        product=cls.normalize(values.get(mapped["product"], "")) or cls.normalize(values.get(mapped["product_code"], "")),
                        product_code=cls.normalize(values.get(mapped["product_code"], "")) or cls.normalize(values.get(mapped["product"], "")),
                        category=cls.normalize(values.get(mapped["category"], "")),
                        pattern=cls.normalize(values.get(mapped["pattern"], "")),
                        qty=cls._positive(values.get(mapped["qty"], "0")),
                        unit_price=purchase_price,
                        purchase_price=purchase_price,
                        sale_price=sale_price,
                        pallet=cls._parse_pallet(values.get(mapped["pallet"])),
                        row_number=row_number,
                    )

                    if not item.product:
                        report.errors.append(f"الصف {row_number}: اسم الصنف فارغ.")
                        continue
                    if not item.invoice_no:
                        report.errors.append(f"الصف {row_number}: رقم الفاتورة فارغ.")
                        continue

                    report.rows.append(item)

                except Exception as exc:
                    get_error_logger("invoice_import").warning(
                        "Skipped an invalid invoice spreadsheet row: %s", exc,
                    )
                    report.errors.append(f"الصف {row_number}: بيانات غير صحيحة: {exc}")

            _, product_errors = cls._resolve_products(report)
            report.errors.extend(product_errors)
            return report

        except Exception as exc:
            get_error_logger("invoice_import").exception("Failed to read invoice spreadsheet")
            report.errors.append(f"فشل في قراءة ملف Excel: {exc}")
            return report
        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    def _product_key(name: str) -> str:
        # Ignore only the characters agreed for invoice imports.
        return "".join(char for char in str(name).casefold()
                       if not char.isspace() and char not in "()[]{}-")

    @classmethod
    def _resolve_products(cls, report: ImportReport):
        from app.services.product_service import product_service
        from app.services.supplier_service import supplier_service
        from app.config import CATEGORIES, CATEGORIES_AR
        suppliers = supplier_service.get_all()
        category_map = dict(zip(CATEGORIES_AR, CATEGORIES))
        category_map.update({item: item for item in CATEGORIES})
        index = {}
        for product in product_service.get_all():
            key = cls._product_key(product.get("name") or "")
            if key:
                index.setdefault(key, []).append(product)
        resolved, errors = {}, []
        for fallback_number, row in enumerate(report.rows, start=2):
            number = row.row_number or fallback_number
            name = row.product or row.product_code
            key = cls._product_key(name)
            candidates = list(index.get(key, [])) if key else []
            if row.pattern:
                candidates = [p for p in candidates if product_service._normalize_name(p.get("pattern") or "") == product_service._normalize_name(row.pattern)]
            if row.category:
                candidates = [p for p in candidates if p["category"] == category_map.get(row.category)]
            try:
                supplier = supplier_service.resolve_reference(row.supplier, suppliers)
            except Exception as exc:
                errors.append(f"الصف {number}: {exc}")
                continue
            if supplier:
                candidates = [p for p in candidates if p["supplier_id"] == supplier["id"]]
            if not key:
                errors.append(f"الصف {number}: اسم الصنف فارغ أو لا يحتوي على حروف أو أرقام.")
            elif not candidates:
                errors.append(f"الصف {number}: الصنف «{name}» غير موجود في قاعدة البيانات.")
            elif len(candidates) > 1:
                names = "، ".join(product["name"] for product in candidates)
                errors.append(f"الصف {number}: الاسم «{name}» يطابق أكثر من صنف بعد تجاهل المسافات والأقواس والشرطة: {names}.")
            else:
                resolved[id(row)] = candidates[0]
        return resolved, errors

    @classmethod
    @log_failures("save imported invoices")
    def save_report_to_db(cls, report: ImportReport) -> int:
        from app.config import CATEGORIES, CATEGORIES_AR
        from app.db import db
        from app.services.invoice_service import invoice_service
        from app.services.supplier_service import supplier_service

        if report.errors:
            raise ValueError("لم يتم استيراد أي فاتورة:\n" + "\n".join(report.errors))
        created_invoices = 0
        # Validation and all invoice writes share one transaction. No product or
        # supplier is created as a side effect of matching spreadsheet names.
        with db.transaction(immediate=True):
            resolved, errors = cls._resolve_products(report)
            category_map = dict(zip(CATEGORIES_AR, CATEGORIES))
            category_map.update({item: item for item in CATEGORIES})
            grouped = {}
            for fallback_number, row in enumerate(report.rows, start=2):
                number = row.row_number or fallback_number
                if not row.invoice_no:
                    errors.append(f"الصف {number}: رقم الفاتورة فارغ.")
                product = resolved.get(id(row))
                if product and row.category:
                    category = category_map.get(row.category)
                    if category is None:
                        errors.append(f"الصف {number}: تصنيف الصنف غير معروف: {row.category}.")
                    elif category != product["category"]:
                        errors.append(f"الصف {number}: تصنيف الصنف «{row.product}» لا يطابق التصنيف المسجل.")
                if product:
                    grouped.setdefault((product["supplier_id"], row.invoice_no), []).append(row)
            plans = []
            suppliers = supplier_service.get_all()
            for (_, invoice_no), rows in grouped.items():
                first = rows[0]
                products = [resolved.get(id(row)) for row in rows]
                if any(product is None for product in products):
                    continue
                supplier_ids = {product["supplier_id"] for product in products}
                supplier = supplier_service.resolve_reference(first.supplier, suppliers)
                if len(supplier_ids) != 1:
                    for row, product in zip(rows, products):
                        errors.append(f"الصف {row.row_number or report.rows.index(row) + 2}: الصنف «{row.product}» تابع للمورد «{product['supplier_name']}»؛ أصناف الفاتورة ليست لمورد واحد.")
                    continue
                supplier_id = next(iter(supplier_ids))
                if supplier and supplier["id"] != supplier_id:
                    for row in rows:
                        errors.append(f"الصف {row.row_number or report.rows.index(row) + 2}: الصنف «{row.product}» تابع لمورد آخر. صحح المورد في Excel.")
                    continue
                items = [{
                    "product_id": product["id"],
                    "quantity": float(row.qty or 0),
                    "pallet": cls._parse_pallet(row.pallet),
                    "purchase_usd": float(row.purchase_price or row.unit_price or 0),
                    "sale_dzd": float(row.sale_price or row.unit_price or 0),
                } for row, product in zip(rows, products)]
                dates = {cls._parse_date(row.date) for row in rows}
                if len(dates) != 1:
                    errors.append(f"الفاتورة {invoice_no}: تواريخ الصفوف مختلفة.")
                plans.append((invoice_no, supplier_id, cls._parse_date(first.date), items))
            if errors:
                raise ValueError("لم يتم استيراد أي فاتورة:\n" + "\n".join(errors))
            for invoice_no, supplier_id, invoice_date, items in plans:
                if db.query_one("SELECT id FROM invoices WHERE supplier_invoice_no = ? AND supplier_id = ?", (invoice_no, supplier_id)):
                    continue
                invoice_service.save_draft(
                    invoice_id=None, supplier_id=supplier_id,
                    supplier_invoice_no=invoice_no, invoice_date=invoice_date,
                    items=items, payments=[], discounts=[], expenses=[], display_fields=None,
                )
                created_invoices += 1
        return created_invoices
