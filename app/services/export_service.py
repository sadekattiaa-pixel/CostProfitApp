"""
خدمة التصدير — PDF و Excel للتقارير والفواتير.
"""
import os
import re
from decimal import Decimal, ROUND_HALF_UP
from xml.sax.saxutils import escape
from reportlab.pdfgen import canvas
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.fonts import tt2ps
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, PageBreak, Flowable,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

import arabic_reshaper
from bidi.algorithm import get_display

from app.services.settings_service import settings
from app.services.error_logging import log_failures
from app.paths import get_company_dir, get_products_images_dir


# ============================================================
# تسجيل خطوط عربية
# ============================================================
_ARABIC_FONT_NAME = "Helvetica"
_FONT_REGISTERED = False


def _try_register_arabic_font():
    global _ARABIC_FONT_NAME, _FONT_REGISTERED
    if _FONT_REGISTERED:
        return

    candidates = [
        ("Cairo", Path("C:/Windows/Fonts/Cairo-Regular.ttf")),
        ("Cairo", Path("C:/Windows/Fonts/cairo.ttf")),
        ("Amiri", Path("C:/Windows/Fonts/Amiri-Regular.ttf")),
        ("Arial", Path("C:/Windows/Fonts/arial.ttf")),
        ("Tahoma", Path("C:/Windows/Fonts/tahoma.ttf")),
    ]
    app_root = Path(__file__).resolve().parent.parent.parent
    assets_fonts = app_root / "assets" / "fonts"
    if assets_fonts.exists():
        for f in assets_fonts.glob("*.ttf"):
            candidates.insert(0, (f.stem, f))

    for name, path in candidates:
        if path.exists():
            try:
                pdfmetrics.registerFont(TTFont(name, str(path)))
                bold_name = name
                bold_candidates = [
                    path.with_name(path.stem.replace("Regular", "Bold") + path.suffix),
                    path.with_name(path.stem + "-Bold" + path.suffix),
                    path.with_name(path.stem + "bd" + path.suffix),
                ]
                for bold_path in bold_candidates:
                    if bold_path != path and bold_path.exists():
                        try:
                            bold_name = name + "-Bold"
                            pdfmetrics.registerFont(TTFont(bold_name, str(bold_path)))
                            break
                        except Exception:
                            bold_name = name
                pdfmetrics.registerFontFamily(name, normal=name, bold=bold_name,
                                              italic=name, boldItalic=bold_name)
                _ARABIC_FONT_NAME = name
                _FONT_REGISTERED = True
                return
            except Exception:
                continue
    _FONT_REGISTERED = True


def _ar(text) -> str:
    if text is None:
        return ""
    try:
        return get_display(arabic_reshaper.reshape(str(text)))
    except Exception:
        return str(text)


def _fmt_money(v) -> str:
    try:
        return f"{float(v):,.2f}"
    except (TypeError, ValueError):
        return "0.00"


def _pdf_number(value, decimals=0) -> str:
    try:
        number = Decimal(str(value)).quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP)
        return f"{number:,.{decimals}f}"
    except Exception:
        return "—" if value is None or value == "" else str(value)


def _pdf_text(value, keep_decimals=False) -> str:
    text = str(value)
    match = re.fullmatch(r"([+$\s]*)(-?\d[\d,]*(?:\.\d+)?)(\s*(?:%|د\.ل|LYD|USD|\$)?)", text)
    if not match:
        return text
    return match[1] + _pdf_number(match[2].replace(",", ""), 2 if keep_decimals else 0) + match[3]


class _SingleLineProduct(Flowable):
    """Product identity retains its font size and never wraps."""
    def __init__(self, text, font_name):
        super().__init__()
        self.text, self.font_name = _ar(text), font_name
        self.width = pdfmetrics.stringWidth(self.text, font_name, 8)
        self.height = 9

    def wrap(self, available_width, available_height):
        return self.width, self.height

    def draw(self):
        self.canv.setFont(self.font_name, 8)
        self.canv.drawString(0, 2, self.text)


class _CompactTableHeading(_SingleLineProduct):
    """Keep a narrow heading within a fixed 14-point row at 8 points."""
    def __init__(self, text, font_name, available_width):
        super().__init__(text, font_name)
        self.lines = None
        if self.width > available_width:
            candidates = [(self.text[:index].strip(), self.text[index:].strip()) for index in range(1, len(self.text))]
            candidates = [parts for parts in candidates if max(pdfmetrics.stringWidth(part, font_name, 8) for part in parts) <= available_width]
            if candidates:
                self.lines = min(candidates, key=lambda parts: (abs(len(parts[0]) - len(parts[1])), abs(pdfmetrics.stringWidth(parts[0], font_name, 8) - pdfmetrics.stringWidth(parts[1], font_name, 8))))
                self.width = max(pdfmetrics.stringWidth(part, font_name, 8) for part in self.lines)
                self.height = 12

    def draw(self):
        if not self.lines:
            return super().draw()
        self.canv.setFont(self.font_name, 8)
        for line, baseline in zip(self.lines, (7, 1)):
            self.canv.drawCentredString(self.width / 2, baseline, line)


class _InvoiceMarks(Flowable):
    """Independently positioned images; signature is drawn after the stamp."""
    def __init__(self, marks, page_width):
        super().__init__()
        from PIL import Image as PillowImage
        from app.paths import get_images_dir
        self.width = page_width
        self.entries = []
        for kind in ("stamp", "signature"):
            filename = marks.get("invoice_" + kind + "_image")
            if not filename:continue
            root = get_images_dir().resolve()
            path = (root / "Marks" / filename).resolve()
            if root not in path.parents or not path.is_file():continue
            width = float(marks.get("invoice_" + kind + "_width_mm") or (30 if kind == "stamp" else 40)) * mm
            x = float(marks.get("invoice_" + kind + "_x_mm") or 0) * mm
            y = float(marks.get("invoice_" + kind + "_y_mm") or 0) * mm
            with PillowImage.open(path) as image:height = width * image.height / image.width
            if width <= 0 or x < 0 or y < 0 or x + width > page_width + 0.1:
                raise ValueError("موقع الختم أو التوقيع خارج عرض الصفحة؛ عدّل العرض والمسافة من اليسار. / Stamp or signature exceeds printable width.")
            self.entries.append((str(path), x, y, width, height))
        self.height = max((y + height for _, x, y, width, height in self.entries), default=0)

    def draw(self):
        for path, x, y, width, height in self.entries:
            self.canv.drawImage(path, x, self.height - y - height, width=width, height=height, mask="auto")


def _current_marks():
    return {"invoice_" + kind + "_" + field: settings.get("invoice_" + kind + "_" + field, default) for kind, defaults in (("stamp", (30, 130, 5)), ("signature", (40, 140, 5))) for field, default in zip(("image", "width_mm", "x_mm", "y_mm"), ("",) + defaults)}


def _fixed_product_width(font_name):
    """Freeze a catalog-based width shared by all invoice exports."""
    from app.db import db
    saved = settings.get("invoice_pdf_product_width_v1")
    if saved:
        return float(saved)
    products = db.query_all("SELECT CASE WHEN category = 'Batteries' THEN code ELSE COALESCE(NULLIF(name, ''), code) END AS identity FROM products")
    width = max([110.0] + [pdfmetrics.stringWidth(_ar(p["identity"] or ""), font_name, 10) + 8 for p in products])
    if products:
        settings.set("invoice_pdf_product_width_v1", width)
    return width


def _auto_pdf_widths(rows, available, font_name, font_size, image_columns=()):
    """Measure text and distribute page width, allowing long labels to wrap."""
    count = len(rows[0])
    minimum, preferred = [], []
    for index in range(count):
        texts = [str(row[index]) for row in rows]
        token = max((pdfmetrics.stringWidth(_ar(word), font_name, font_size)
                     for text in texts for word in text.split()), default=12)
        line = max((pdfmetrics.stringWidth(_ar(text), font_name, font_size) for text in texts), default=12)
        floor = min(max(token + 6, 24), available * 0.22)
        want = min(max(line + 8, floor), available * 0.35)
        if index in image_columns:
            floor = want = 32
        minimum.append(floor)
        preferred.append(want)
    if sum(minimum) >= available:
        return [width * available / sum(minimum) for width in minimum]
    slack = available - sum(minimum)
    growth = [max(0, want-floor) for want, floor in zip(preferred, minimum)]
    if sum(growth) > slack:
        return [floor + slack * extra / sum(growth) for floor, extra in zip(minimum, growth)]
    spare = available - sum(preferred)
    return [width + spare * width / sum(preferred) for width in preferred]


# ============================================================
# ReportData
# ============================================================
class ReportData:
    def __init__(self, title, headers, rows, kpis=None, subtitle=""):
        self.title = title
        self.subtitle = subtitle
        self.headers = headers
        self.rows = rows
        self.kpis = kpis or []


# ============================================================
# الإعدادات الافتراضية
# ============================================================
DEFAULT_INVOICE_SECTIONS = {
    "header":    True,
    "info":      True,
    "items":     True,
    "payments":  True,
    "discounts": True,
    "expenses":  True,
    "summary":   True,
    "footer":    True,
    "marks":     True,
}


INVOICE_ITEM_COLUMNS = [
    ("product_name",  "Product", "Product",             True),
    ("pattern",       "Pattern",      "النقشة",             True),
    ("image",         "layout",        "layout",            False),
    ("category",      "Category",     "التصنيف",           False),
    ("ah",            "AH",           "AH",                False),
    ("l",             "L",            "L",                 False),
    ("w",             "W",            "W",                 False),
    ("h",             "H",            "H",                 False),
    ("oh",            "OH",           "OH",                False),
    ("layout",        "layout",       "layout",           False),
    ("pallet",        "Pallet",       "Pallet",           True),
    ("quantity",      "Qty",          "Qty",            True),
    ("purchase_usd",  "Price $",   "الشراء $",          True),
    ("purchase_dzd",  "Unit LYD",     "الوحدة د.ل",        True),
    ("sale_dzd",      "Sale LYD",     "البيع د.ل",         True),
    ("total_usd",     "Total $",      "الإجمالي $",        True),
    ("total_dzd",     "Total LYD",    "الإجمالي د.ل",      True),
    ("total_sale",    "Total Sale",   "إجمالي البيع",      True),
    ("unit_profit",   "Unit Profit",  "ربح الوحدة",        False),
    ("total_profit",  "Profit",       "الربح",             True),
]

INVOICE_ITEM_COLUMN_KEYS = [c[0] for c in INVOICE_ITEM_COLUMNS]
DEFAULT_INVOICE_ITEM_COLUMNS = [c[0] for c in INVOICE_ITEM_COLUMNS if c[3]]


SUPPLIER_FIELDS = [
    ("name",          "Name",             "الاسم",              True),
    ("code",          "Code",             "الكود",              False),
    ("address",       "Address",          "العنوان",            True),
    ("phone",         "Phone",            "الهاتف",             True),
    ("email",         "Email",            "البريد",             False),
    ("invoice_no",    "Supplier Inv No",  "رقم فاتورة المورد",  True),
]

SUPPLIER_FIELD_KEYS = [f[0] for f in SUPPLIER_FIELDS]
DEFAULT_SUPPLIER_FIELDS = [f[0] for f in SUPPLIER_FIELDS if f[3]]


# ============================================================
# خدمة التصدير
# ============================================================

class _InvoiceNumberedCanvas(canvas.Canvas):
    """Write invoice identification and total page count on every page."""
    def __init__(self, *args, invoice_no="", font_name="Helvetica", is_ar=False, **kwargs):
        super().__init__(*args, **kwargs)
        self._page_states = []
        self._invoice_no = invoice_no
        self._footer_font = font_name
        self._is_ar = is_ar

    def showPage(self):
        self._page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._page_states)
        for state in self._page_states:
            self.__dict__.update(state)
            width, _ = self._pagesize
            self.setFont(self._footer_font, 6)
            self.setFillColor(colors.HexColor("#64748B"))
            page_text = f"صفحة {self._pageNumber} من {total}" if self._is_ar else f"Page {self._pageNumber} of {total}"
            self.drawString(10 * mm, 8 * mm, _ar(page_text))
            self.drawRightString(width - 10 * mm, 8 * mm, _ar(self._invoice_no))
            super().showPage()
        super().save()


def safe_filename(value: str) -> str:
    import re
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value)).strip().rstrip('. ')
    if name.split('.')[0].upper() in {'CON', 'PRN', 'AUX', 'NUL'} or name.split('.')[0].upper() in {f'{p}{n}' for p in ('COM','LPT') for n in range(1,10)}:
        name = '_' + name
    return name[:150] or 'Invoice'

def _sheet_title(value) -> str:
    import re
    return re.sub(r'[\[\]:*?/\\\x00-\x1f]', '_', str(value))[:31].strip("'") or 'Report'

def _save_literal_workbook(workbook, file_path):
    # These exports contain values, never executable spreadsheet formulas.
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
    workbook.save(file_path)


class ExportService:

    @staticmethod
    def _category_columns(invoice, columns):
        items = invoice.get("items") or []
        tires = [item for item in items if item.get("product_category") in ("Car Tires", "Truck Tires")]
        if not tires:
            columns = [key for key in columns if key != "pattern"]
        elif len(tires) == len(items):
            columns = [key for key in columns if key not in ("ah", "l", "w", "h", "oh", "layout", "image")]
        return columns

    @log_failures("export products to Excel")
    def export_products_excel(
        self,
        products: list[dict],
        file_path: str,
        language: str = "en",
    ) -> str:
        from app.config import CATEGORIES, CATEGORIES_AR

        is_ar = language == "ar"
        headers = (
            ["Product", "التصنيف", "المورد", "AH", "L", "W", "H", "OH", "layout", "Pallet", "layout"]
            if is_ar else
            ["Product", "Category", "Supplier", "AH", "L", "W", "H", "OH", "layout", "Pallet", "layout"]
        )
        has_tires = any(product.get("category") != "Batteries" for product in products)
        tires_only = bool(products) and all(product.get("category") != "Batteries" for product in products)
        selected = [0, 1, 2] if tires_only else [0, 1, 2, 3, 4, 5, 6, 7, 8, 10, 9]
        headers = [headers[index] for index in selected]
        if has_tires:
            headers[0] = "Product"
            headers.insert(1, "النقشة" if is_ar else "Pattern")
        wb = Workbook()
        ws = wb.active
        ws.title = "الأصناف" if is_ar else "Products"
        ws.sheet_view.rightToLeft = is_ar
        ws.freeze_panes = "A2"

        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        border_side = Side(style="thin", color="D8E1EB")
        border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_font = Font(bold=True, color="FFFFFF", size=11, name="Cairo")

        for column, label in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=column, value=label)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border

        if headers.count("layout") == 2:
            left = headers.index("layout") + 1
            ws.merge_cells(start_row=1, start_column=left, end_row=1, end_column=left + 1)
        for row_index, product in enumerate(products, start=2):
            category = product.get("category") or ""
            if is_ar and category in CATEGORIES:
                category = CATEGORIES_AR[CATEGORIES.index(category)]
            values = [
                product.get("code", ""),
                category,
                product.get("supplier_name", ""),
                product.get("ah"),
                product.get("l"),
                product.get("w"),
                product.get("h"),
                product.get("oh"),
                product.get("layout"),
                product.get("pallet"),
                None,
            ]
            values = [values[index] for index in selected]
            if has_tires:
                values.insert(1, product.get("pattern") or "")
            for column, value in enumerate(values, start=1):
                cell = ws.cell(row=row_index, column=column, value=value)
                cell.alignment = center
                cell.border = border

            image_path = self._get_product_image_path(product)
            if image_path and 10 in selected:
                picture = ExcelImage(str(image_path))
                scale = min(48 / picture.width, 40 / picture.height)
                picture.width *= scale
                picture.height *= scale
                picture.anchor = f"{get_column_letter(selected.index(10) + 1 + int(has_tires))}{row_index}"
                ws.add_image(picture)
                ws.row_dimensions[row_index].height = 32

        widths = [28, 20, 24, 10, 10, 10, 10, 10, 13, 13, 12]
        widths = [widths[index] for index in selected]
        if has_tires:
            widths.insert(1, 22)
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(products) + 1)}"
        _save_literal_workbook(wb, file_path)
        return file_path

    # --------------------------------------------------------
    # Excel للتقارير
    # --------------------------------------------------------
    @log_failures("export a report to Excel")
    def export_excel(self, report: ReportData, file_path: str) -> str:
        wb = Workbook()
        ws = wb.active
        ws.title = _sheet_title(report.title)
        ws.sheet_view.rightToLeft = True

        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        kpi_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12, name="Cairo")
        title_font = Font(bold=True, size=16, name="Cairo")
        kpi_font = Font(bold=True, size=12, name="Cairo")
        border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        row_idx = 1
        num_cols = max(len(report.headers), 2)

        ws.merge_cells(start_row=row_idx, start_column=1,
                       end_row=row_idx, end_column=num_cols)
        c = ws.cell(row=row_idx, column=1, value=report.title)
        c.font = title_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        row_idx += 1

        ws.merge_cells(start_row=row_idx, start_column=1,
                       end_row=row_idx, end_column=num_cols)
        c = ws.cell(row=row_idx, column=1,
                    value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        c.alignment = Alignment(horizontal="center")
        c.font = Font(size=10, name="Cairo", color="64748B")
        row_idx += 2

        if report.kpis:
            for kpi_label, kpi_value in report.kpis:
                cell_l = ws.cell(row=row_idx, column=1, value=kpi_label)
                cell_l.font = kpi_font
                cell_l.fill = kpi_fill
                cell_l.border = border
                cell_v = ws.cell(row=row_idx, column=2, value=kpi_value)
                cell_v.font = kpi_font
                cell_v.fill = kpi_fill
                cell_v.border = border
                row_idx += 1
            row_idx += 1

        for col_idx, h in enumerate(report.headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        row_idx += 1

        for row in report.rows:
            for col_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border
            row_idx += 1

        for col_idx in range(1, len(report.headers) + 1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = 20

        _save_literal_workbook(wb, file_path)
        return file_path

    # --------------------------------------------------------
    # PDF للتقارير
    # --------------------------------------------------------
    @log_failures("export a report to PDF")
    def export_pdf(self, report: ReportData, file_path: str) -> str:
        _try_register_arabic_font()
        font_name = _ARABIC_FONT_NAME

        doc = SimpleDocTemplate(
            file_path,
            pagesize=landscape(A4),
            rightMargin=15 * mm,
            leftMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm,
            title=report.title,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("T", parent=styles["Title"], fontName=font_name,
                                     fontSize=20, alignment=1,
                                     textColor=colors.HexColor("#0A1220"))
        sub_style = ParagraphStyle("S", parent=styles["Normal"], fontName=font_name,
                                   fontSize=11, alignment=1,
                                   textColor=colors.HexColor("#64748B"))
        h_style = ParagraphStyle("H", parent=styles["Normal"], fontName=font_name,
                                 fontSize=9, leading=11, textColor=colors.black, alignment=1)
        c_style = ParagraphStyle("C", parent=styles["Normal"], fontName=font_name,
                                 fontSize=9, leading=11, alignment=1)
        kpi_lbl_style = ParagraphStyle("KL", parent=styles["Normal"], fontName=font_name,
                                       fontSize=10, textColor=colors.black,
                                       alignment=1)
        kpi_val_style = ParagraphStyle("KV", parent=styles["Normal"], fontName=font_name,
                                       fontSize=12, textColor=colors.black,
                                       alignment=1)
        company_style = ParagraphStyle("Co", parent=styles["Normal"], fontName=font_name,
                                       fontSize=11, alignment=0,
                                       textColor=colors.HexColor("#0A1220"), leading=16)

        elements = []

        company_name = settings.get("company_name") or ""
        company_phone = settings.get("company_phone") or ""
        company_email = settings.get("company_email") or ""
        company_addr = settings.get("company_address") or ""
        logo_filename = settings.get("company_logo") or ""

        logo_path = None
        if logo_filename:
            try:
                candidate = get_company_dir() / logo_filename
                if candidate.exists():
                    logo_path = str(candidate)
            except Exception:
                pass

        if logo_path:
            try:
                logo_img = Image(logo_path, width=24 * mm, height=24 * mm)
            except Exception:
                logo_img = Paragraph(" ", sub_style)
        else:
            logo_img = Paragraph(" ", sub_style)

        lines = []
        if company_name:
            lines.append(f"<b>{escape(_ar(company_name))}</b>")
        if company_phone:
            lines.append(escape(_ar(f"هاتف: {company_phone}")))
        if company_email:
            lines.append(escape(_ar(f"البريد: {company_email}")))
        if company_addr:
            lines.append(escape(_ar(f"العنوان: {company_addr}")))

        company_html = "<br/>".join(lines) if lines else " "
        company_para = Paragraph(company_html, company_style)

        ht = Table([[logo_img, company_para]],
                   colWidths=[30 * mm, doc.width - 30 * mm])
        ht.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (0, 0), "LEFT"),
            ("ALIGN", (1, 0), (1, 0), "RIGHT"),
            ("LEFTPADDING", (0, 0), (0, 0), 0),
            ("RIGHTPADDING", (1, 0), (1, 0), 0),
        ]))
        elements.append(ht)
        elements.append(Spacer(1, 6))

        sep = Table([[""]], colWidths=[doc.width], rowHeights=[1])
        sep.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#2563EB")),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        elements.append(sep)
        elements.append(Spacer(1, 12))

        elements.append(Paragraph(escape(_ar(report.title)), title_style))
        elements.append(Spacer(1, 4))
        date_str = f"تاريخ التقرير: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        elements.append(Paragraph(_ar(date_str), sub_style))
        elements.append(Spacer(1, 12))

        if report.kpis:
            rows_kpi = []
            current = []
            for kpi in report.kpis:
                current.append(kpi)
                if len(current) == 4:
                    rows_kpi.append(current); current = []
            if current:
                while len(current) < 4:
                    current.append(("", ""))
                rows_kpi.append(current)

            for row in rows_kpi:
                tdata = [
                    [Paragraph(escape(_ar(lbl)), kpi_lbl_style) for lbl, _ in row],
                    [Paragraph(escape(_ar(_pdf_text(val, "$" in str(val) or "USD" in str(val)))), kpi_val_style) for _, val in row],
                ]
                t = Table(tdata, colWidths=[50 * mm] * 4)
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                    ("BOX", (0, 0), (-1, -1), 1.2, colors.black),
                    ("INNERGRID", (0, 0), (-1, -1), 1.2, colors.black),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]))
                elements.append(t)
                elements.append(Spacer(1, 6))
            elements.append(Spacer(1, 10))

        if report.rows:
            numeric_terms = ("$", "USD", "LYD", "د.ل", "qty", "quantity", "profit", "margin", "change", "count", "sales", "sale", "cost", "purchase", "payment", "discount", "expense", "rate",
                             "Qty", "القيمة", "العدد", "الربح", "هامش", "تغير", "تكلفة", "البيع", "الشراء", "مدفوع", "خصم", "مصروف")
            formatted = [list(report.headers)]
            for row in report.rows:
                formatted.append([
                    _pdf_text(value, "$" in str(header) or "USD" in str(header))
                    if any(term.casefold() in str(header).casefold() for term in numeric_terms) else str(value)
                    for header, value in zip(report.headers, row)
                ])
            data = [[Paragraph(escape(_ar(value)), h_style if index == 0 else c_style)
                     for value in row] for index, row in enumerate(formatted)]
            widths = _auto_pdf_widths(formatted, doc.width, font_name, 9)
            table = Table(data, colWidths=widths, repeatRows=1, splitByRow=1, splitInRow=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 1.2, colors.black),
                
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]))
            elements.append(table)
        else:
            elements.append(Paragraph(_ar("لا توجد بيانات"), sub_style))

        doc.build(elements)
        return file_path

    # --------------------------------------------------------
    # PDF للفاتورة
    # --------------------------------------------------------
    @log_failures("export an invoice to PDF")
    def export_invoice_pdf(self, invoice: dict, file_path: str,
                           sections: dict | None = None,
                           columns: list[str] | None = None,
                           supplier_fields: list[str] | None = None) -> str:
        if invoice.get("status") != "archived":
            invoice = dict(invoice, notes=settings.get("invoice_notes") or "", notes2=settings.get("invoice_notes2") or "")
        sections = dict(DEFAULT_INVOICE_SECTIONS) if sections is None else sections
        columns = list(DEFAULT_INVOICE_ITEM_COLUMNS) if columns is None else list(columns)
        columns = [key for key in INVOICE_ITEM_COLUMN_KEYS if key in columns]
        if "product_name" not in columns:
            columns.insert(0, "product_name")
        columns = self._category_columns(invoice, self._order_pdf_invoice_columns(columns))
        supplier_fields = list(DEFAULT_SUPPLIER_FIELDS) if supplier_fields is None else supplier_fields
        supplier_fields = [key for key in SUPPLIER_FIELD_KEYS if key in supplier_fields]
        _try_register_arabic_font()
        font_name = _ARABIC_FONT_NAME
        table_font = tt2ps(font_name, 1, 0)
        is_ar = self._is_arabic()
        label = lambda ar, en: ar if is_ar else en
        reserved = {"product_name": _fixed_product_width(font_name), "quantity": 40.0,
                    "image": 10.0 + 4, "layout": 34.0}
        minimum_width = sum(reserved.get(key, max([24.0] + [
            pdfmetrics.stringWidth(_ar(self._get_item_column_value(item, key, pdf=True)), font_name, 8) + 6
            for item in invoice.get("items", [])])) for key in columns)
        page_size = landscape(A4) if minimum_width > A4[0] - 20 * mm else A4
        doc = SimpleDocTemplate(file_path, pagesize=page_size, rightMargin=10 * mm,
                                leftMargin=10 * mm, topMargin=10 * mm,
                                bottomMargin=16 * mm, title=invoice.get("invoice_no", "Invoice"))
        styles = getSampleStyleSheet()
        invoice_header_style = ParagraphStyle("IInv", parent=styles["Normal"],
            fontName=font_name, fontSize=7, alignment=2, leading=9)
        sub_style = ParagraphStyle("IS", parent=styles["Normal"], fontName=font_name,
            fontSize=8, alignment=1, textColor=colors.HexColor("#64748B"))
        section_style = ParagraphStyle("ISec", parent=styles["Normal"], fontName=font_name,
            fontSize=10, leading=12, textColor=colors.HexColor("#2563EB"), spaceBefore=6, spaceAfter=6)
        company_style = ParagraphStyle("ICo", parent=styles["Normal"], fontName=font_name,
            fontSize=8, alignment=0, leading=10)
        supplier_style = ParagraphStyle("ISupplier", parent=company_style, fontSize=8, leading=10)
        elements = []
        page_width = doc.width
        show_header = sections.get("header", True)
        show_info = sections.get("info", True)
        if show_header or show_info:
            company_name = (invoice.get("_company_snapshot", {}).get("company_name") if "_company_snapshot" in invoice else settings.get("company_name")) or ""
            company_phone = (invoice.get("_company_snapshot", {}).get("company_phone") if "_company_snapshot" in invoice else settings.get("company_phone")) or ""
            company_website = (invoice.get("_company_snapshot", {}).get("company_website") if "_company_snapshot" in invoice else settings.get("company_website")) or ""
            company_email = (invoice.get("_company_snapshot", {}).get("company_email") if "_company_snapshot" in invoice else settings.get("company_email")) or ""
            company_addr = (invoice.get("_company_snapshot", {}).get("company_address") if "_company_snapshot" in invoice else settings.get("company_address")) or ""
            logo_filename = (invoice.get("_company_snapshot", {}).get("company_logo") if "_company_snapshot" in invoice else settings.get("company_logo")) or ""

            logo_path = None
            if logo_filename:
                try:
                    candidate = get_company_dir() / logo_filename
                    if candidate.exists():
                        logo_path = str(candidate)
                except Exception:
                    pass

            if logo_path:
                try:
                    logo_img = Image(logo_path, width=16 * mm, height=16 * mm)
                except Exception:
                    logo_img = Paragraph(" ", sub_style)
            else:
                logo_img = Paragraph(" ", sub_style)

            left_width = page_width * 0.62
            header_rows = []

            if show_header:
                company_lines = []
                if company_name:
                    company_lines.append("<b>" + escape(_ar(company_name)) + "</b>")
                for value in (company_addr, company_phone, company_email, company_website):
                    if value:
                        company_lines.append("<b>" + escape(_ar(value)) + "</b>")

                company_para = Paragraph(
                    "<br/>".join(company_lines) if company_lines else " ",
                    company_style,
                )
                company_block = Table(
                    [[logo_img, company_para]],
                    colWidths=[20 * mm, left_width - 20 * mm],
                )
                company_block.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ]))
                header_rows.append([company_block, Paragraph(" ", sub_style)])

            supplier_para = Paragraph(" ", supplier_style)
            if show_info and supplier_fields:
                supplier_values = {
                    "name": invoice.get("supplier_name") or "",
                    "address": invoice.get("supplier_address") or "",
                    "phone": invoice.get("supplier_phone") or "",
                    "email": invoice.get("supplier_email") or "",
                }
                selected_fields = set(supplier_fields)
                supplier_order = ("name", "address", "phone", "email")
                supplier_lines = []
                for key in supplier_order:
                    value = supplier_values[key]
                    if key in selected_fields and value:
                        formatted = _ar(value)
                        supplier_lines.append(
                            escape(formatted)
                        )
                if supplier_lines:
                    supplier_para = Paragraph("<br/>".join(supplier_lines), supplier_style)

            invoice_lines = []
            if show_info:
                inv_no = invoice.get("invoice_no", "")
                date_str = invoice.get("invoice_date", "")
                supplier_inv_no = invoice.get("supplier_invoice_no", "")
                if inv_no:
                    invoice_lines.append(escape(_ar(inv_no)))
                if date_str:
                    invoice_lines.append(escape(_ar(date_str)))
                if supplier_inv_no:
                    invoice_lines.append(escape(_ar(supplier_inv_no)))
            invoice_para = Paragraph(
                "<br/>".join(invoice_lines) if invoice_lines else " ",
                invoice_header_style,
            )

            info_row_index = None
            if show_info:
                # Supplier and invoice data share the same row below the company header.
                info_row_index = len(header_rows)
                header_rows.append([supplier_para, invoice_para])

            ht = Table(
                header_rows,
                colWidths=[left_width, page_width - left_width],
            )
            header_commands = [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (0, 0), (0, -1), "LEFT"),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
            if info_row_index is not None and show_header:
                header_commands.append(
                    ("TOPPADDING", (0, info_row_index), (-1, info_row_index), 4)
                )
            ht.setStyle(TableStyle(header_commands))
            elements.append(ht)

            sep = Table([[""]], colWidths=[page_width], rowHeights=[1])
            sep.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#2563EB")),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]))
            elements.append(sep)
            elements.append(Spacer(1, 6))


        notes = "\n\n".join(str(invoice.get(key) or "").strip() for key in ("notes", "notes2") if str(invoice.get(key) or "").strip())
        note_blocks = []
        if sections.get("footer", True) and notes:
            note_style = ParagraphStyle("InvoiceNotes", parent=styles["Normal"],
                fontName=font_name, fontSize=8, leading=10, alignment=0,
                textColor=colors.HexColor("#334155"))
            note_blocks = [Spacer(1, 8), Paragraph(
                "<br/>".join(escape(_ar(line)) if line else "&#160;" for line in notes.splitlines()),
                note_style,
            )]

        if sections.get("items", True):
            elements.append(Paragraph(_ar(label("الأصناف", "Products")), section_style))
            items = invoice.get("items") or []
            if items:
                column_names = {column[0]: column[2] if is_ar else column[1] for column in INVOICE_ITEM_COLUMNS}
                additive_columns = ("quantity", "pallet", "total_usd", "total_dzd", "total_sale", "total_profit")
                total_values = []
                for key in columns:
                    if key not in additive_columns:
                        total_values.append("")
                        continue
                    values = [item.get(key) if key != "pallet" or item.get(key) is not None else item.get("product_pallet") for item in items]
                    total_values.append(_pdf_number(sum(float(value or 0) for value in values), 2 if key == "total_usd" else 0))

                product_width = _fixed_product_width(font_name)

                def make_items_table(font_size):
                    raw = [[column_names[key] for key in columns]] + [
                        ["" if key == "image" else self._get_item_column_value(item, key, pdf=True) for key in columns]
                        for item in items] + [total_values]
                    fixed = {"product_name": product_width, "quantity": 40.0, "image": 10.0 + 4, "layout": 34.0}
                    flexible = [i for i, key in enumerate(columns) if key not in fixed]
                    widths = [fixed.get(key, 0) for key in columns]
                    if flexible:
                        remaining = page_width - sum(widths)
                        measured = _auto_pdf_widths([[row[i] for i in flexible] for row in raw], remaining, font_name, font_size)
                        for index, width in zip(flexible, measured):
                            widths[index] = width
                        # Measure against the previous table before narrowing the
                        # five specification columns to exactly half their width.
                        previous_fixed = dict(fixed, image=7 * mm + 4)
                        previous_remaining = page_width - sum(previous_fixed.get(key, 0) for key in columns)
                        previous = _auto_pdf_widths([[row[i] for i in flexible] for row in raw], previous_remaining, font_name, font_size)
                        narrow = {"ah", "l", "w", "h", "oh"}
                        released = 0.0
                        for index, old_width in zip(flexible, previous):
                            if columns[index] in narrow:
                                released += widths[index] - old_width / 2
                                widths[index] = old_width / 2
                        other = [index for index in flexible if columns[index] not in narrow]
                        total_other = sum(widths[index] for index in other)
                        for index in other:
                            widths[index] += released * widths[index] / total_other
                    if "oh" in columns:
                        oh_width = widths[columns.index("oh")]
                    else:
                        # Derive OH's width as if that optional column were visible.
                        virtual_columns = columns + ["oh"]
                        virtual_raw = [raw[0] + ["OH"]] + [raw[index + 1] + [self._get_item_column_value(item, "oh", pdf=True)] for index, item in enumerate(items)] + [raw[-1] + [""]]
                        virtual_flexible = [i for i, key in enumerate(virtual_columns) if key not in fixed]
                        virtual_remaining = page_width - sum(dict(fixed, image=7 * mm + 4).get(key, 0) for key in virtual_columns)
                        virtual_widths = _auto_pdf_widths([[row[i] for i in virtual_flexible] for row in virtual_raw], virtual_remaining, font_name, font_size)
                        oh_width = virtual_widths[-1] / 2
                    shared_columns = {"layout", "image", "pallet"}
                    released = 0.0
                    for index, key in enumerate(columns):
                        if key in shared_columns:
                            released += widths[index] - oh_width
                            widths[index] = oh_width
                    other = [index for index, key in enumerate(columns) if key not in fixed and key not in shared_columns and key not in ("ah", "l", "w", "h", "oh")]
                    total_other = sum(widths[index] for index in other)
                    for index in other:
                        widths[index] += released * widths[index] / total_other
                    data = [[_CompactTableHeading(column_names[key], table_font, widths[index] - (0 if key in ("ah", "l", "w", "h", "oh", "layout", "image", "pallet") else 6)) for index, key in enumerate(columns)]]
                    spans = []
                    if "layout" in columns and "image" in columns:
                        left, right = sorted((columns.index("layout"), columns.index("image")))
                        data[0][left] = _SingleLineProduct("layout", table_font)
                        data[0][right] = ""
                        spans.append(("SPAN", (left, 0), (right, 0)))
                    for item in items:
                        cells = []
                        for key in columns:
                            image_path = self._get_product_image_path(item) if key == "image" else None
                            if image_path:
                                picture_width = max(1, widths[columns.index("image")] - 2)
                                picture = Image(str(image_path), width=picture_width, height=12, kind="proportional")
                                picture.hAlign = "CENTER"
                                cells.append(picture)
                            else:
                                value = "—" if key == "image" else self._get_item_column_value(item, key, pdf=True)
                                cells.append(_SingleLineProduct(value, table_font))
                        data.append(cells)
                    data.append([_SingleLineProduct(value, table_font) for value in total_values])
                    table = Table(data, colWidths=widths, rowHeights=[14] * len(data), repeatRows=1, splitByRow=1)
                    table.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                        ("GRID", (0, 0), (-1, -2), 1.2, colors.black),
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("ALIGN", (columns.index("product_name"), 1), (columns.index("product_name"), -1), "LEFT"),
                        ("NOSPLIT", (0, -2), (-1, -1)),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 3),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                        ("TOPPADDING", (0, 0), (-1, -1), 1),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                    ] + spans + [command for index, key in enumerate(columns) if key in ("ah", "l", "w", "h", "oh", "layout", "image", "pallet") for command in (("LEFTPADDING", (index, 0), (index, -1), 0), ("RIGHTPADDING", (index, 0), (index, -1), 0))]))
                    return table

                items_table = None
                for font_size in (8,):
                    candidate = make_items_table(font_size)
                    combined = elements + [candidate] + note_blocks
                    height = sum(flowable.wrap(page_width, doc.height)[1] + flowable.getSpaceBefore() + flowable.getSpaceAfter() for flowable in combined)
                    items_table = candidate
                    if height <= doc.height - 12:
                        break
                elements.append(items_table)
            else:
                elements.append(Paragraph(_ar(label("لا توجد أصناف", "No products")), sub_style))
        elements.extend(note_blocks)
        if sections.get("marks", True):
            marks = invoice.get("_marks_snapshot", {}) if invoice.get("status") == "archived" else _current_marks()
            mark_block = _InvoiceMarks(marks, page_width)
            if mark_block.entries:
                if mark_block.height > doc.height:raise ValueError("حجم الختم أو التوقيع أكبر من مساحة الصفحة. / Stamp or signature is too tall.")
                elements.append(mark_block)

        financial_keys = ("payments", "discounts", "expenses", "summary")
        if any(sections.get(key, True) for key in financial_keys):
            if elements:
                elements.append(PageBreak())
            list_label_style = ParagraphStyle("ListLabel", parent=styles["Normal"], fontName=font_name,
                fontSize=10, leading=12, alignment=2 if is_ar else 0)
            list_value_style = ParagraphStyle("ListValue", parent=list_label_style, alignment=0 if is_ar else 2)

            def add_list(title, rows, total=None):
                elements.append(Paragraph(_ar(title), section_style))
                data = []
                for text, value, bold in rows:
                    label_text, value_text = escape(_ar(text)), escape(_ar(value))
                    if bold:
                        label_text, value_text = f"<b>{label_text}</b>", f"<b>{value_text}</b>"
                    data.append([Paragraph(label_text, list_label_style), Paragraph(value_text, list_value_style)])
                if total is not None:
                    data.append([
                        Paragraph("<b>" + _ar(label("الإجمالي", "Total")) + "</b>", list_label_style),
                        Paragraph("<b>" + _ar(f"{_pdf_number(total)} د.ل") + "</b>", list_value_style),
                    ])
                if not data:
                    data = [[Paragraph(_ar(label("لا توجد بنود", "No entries")), list_label_style), ""]]
                if is_ar:
                    data = [row[::-1] for row in data]
                table = Table(data, colWidths=[page_width * (0.35 if is_ar else 0.65), page_width * (0.65 if is_ar else 0.35)],
                              hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#E2E8F0")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]))
                elements.extend([table, Spacer(1, 8)])

            for key, title_ar, title_en in (
                ("payments", "المدفوعات", "Payments"), ("discounts", "الخصومات", "Discounts"),
                ("expenses", "المصروفات", "Expenses"),
            ):
                if sections.get(key, True):
                    amounts = invoice.get(key) or []
                    add_list(label(title_ar, title_en), [
                        (item.get("label") or "—", f'{_pdf_number(item.get("amount") or 0)} د.ل', False)
                        for item in amounts
                    ], sum(float(item.get("amount") or 0) for item in amounts))
            if sections.get("summary", True):
                specs = (
                    ("total_payments", "إجمالي المدفوعات", "Total Payments", "money"),
                    ("total_discounts", "إجمالي الخصومات", "Total Discounts", "money"),
                    ("net_payments", "صافي المدفوعات", "Net Payments", "money"),
                    ("total_expenses", "إجمالي المصروفات", "Total Expenses", "money"),
                    ("net_cost", "صافي التكلفة", "Net Cost", "money"),
                    ("total_usd", "القيمة بالدولار", "Total Value (USD)", "usd"),
                    ("exchange_rate", "سعر الصرف", "Exchange Rate", "rate"),
                    ("total_sales", "إجمالي البيع", "Total Sales", "money"),
                    ("total_profit", "إجمالي الربح", "Total Profit", "money"),
                    ("profit_margin", "هامش الربح", "Profit Margin", "percent"),
                )
                rows = []
                for key, title_ar, title_en, kind in specs:
                    value = float(invoice.get(key) or 0)
                    formatted = _pdf_number(value, 4 if kind == "rate" else 2 if kind == "usd" else 0)
                    formatted = "$ " + formatted if kind == "usd" else formatted + (" %" if kind == "percent" else " د.ل" if kind == "money" else "")
                    rows.append((label(title_ar, title_en), formatted, key in ("net_cost", "total_profit", "total_sales")))
                add_list(label("الملخص والأرباح", "Summary and Profit"), rows)
        if not elements:
            elements.append(Paragraph(_ar(invoice.get("invoice_no") or "Invoice"), company_style))
        doc.build(elements, canvasmaker=lambda *args, **kwargs: _InvoiceNumberedCanvas(
            *args, invoice_no=invoice.get("invoice_no") or "", font_name=font_name,
            is_ar=is_ar, **kwargs,
        ))
        return file_path

    # --------------------------------------------------------
    # Excel للفاتورة
    # --------------------------------------------------------
    @log_failures("export an invoice to Excel")
    def export_invoice_excel(self, invoice: dict, file_path: str,
                             sections: dict | None = None,
                             columns: list[str] | None = None,
                             supplier_fields: list[str] | None = None) -> str:
        if invoice.get("status") != "archived":
            invoice = dict(invoice, notes=settings.get("invoice_notes") or "", notes2=settings.get("invoice_notes2") or "")
        if sections is None:
            sections = dict(DEFAULT_INVOICE_SECTIONS)
        if columns is None:
            columns = list(DEFAULT_INVOICE_ITEM_COLUMNS)
        if supplier_fields is None:
            supplier_fields = list(DEFAULT_SUPPLIER_FIELDS)

        if "product_name" not in columns:
            columns = ["product_name"] + list(columns)
        columns = [k for k in INVOICE_ITEM_COLUMN_KEYS if k in columns]
        columns = self._order_pdf_invoice_columns(self._category_columns(invoice, columns))
        supplier_fields = [k for k in SUPPLIER_FIELD_KEYS if k in supplier_fields]

        is_ar = self._is_arabic()

        wb = Workbook()
        ws = wb.active
        ws.title = _sheet_title(invoice.get("invoice_no") or "Invoice")
        ws.sheet_view.rightToLeft = is_ar

        primary_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        section_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
        info_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        title_font = Font(bold=True, size=16, name="Cairo", color="0A1220")
        section_font = Font(bold=True, size=13, name="Cairo", color="1E40AF")
        header_font = Font(bold=True, size=11, name="Cairo", color="FFFFFF")
        label_font = Font(bold=True, size=10, name="Cairo", color="64748B")
        value_font = Font(size=10, name="Cairo", color="0A1220")
        kpi_font = Font(bold=True, size=11, name="Cairo", color="0A1220")

        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        right = Alignment(horizontal="right", vertical="center", wrap_text=True)

        thin = Side(style="thin", color="E2E8F0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        row = 1
        num_cols = max(len(columns), 4)

        # ترويسة الشركة
        if sections.get("header", True):
            company_name = (invoice.get("_company_snapshot", {}).get("company_name") if "_company_snapshot" in invoice else settings.get("company_name")) or ""
            company_phone = (invoice.get("_company_snapshot", {}).get("company_phone") if "_company_snapshot" in invoice else settings.get("company_phone")) or ""
            company_email = (invoice.get("_company_snapshot", {}).get("company_email") if "_company_snapshot" in invoice else settings.get("company_email")) or ""
            company_addr = (invoice.get("_company_snapshot", {}).get("company_address") if "_company_snapshot" in invoice else settings.get("company_address")) or ""

            if company_name:
                ws.merge_cells(start_row=row, start_column=1,
                               end_row=row, end_column=num_cols)
                c = ws.cell(row=row, column=1, value=company_name)
                c.font = title_font
                c.alignment = center
                row += 1

            info_parts = []
            if company_phone:
                info_parts.append(("هاتف" if is_ar else "Phone", company_phone))
            if company_email:
                info_parts.append(("البريد" if is_ar else "Email", company_email))
            if company_addr:
                info_parts.append(("العنوان" if is_ar else "Address", company_addr))

            for label, val in info_parts:
                ws.merge_cells(start_row=row, start_column=1,
                               end_row=row, end_column=num_cols)
                c = ws.cell(row=row, column=1, value=f"{label}: {val}")
                c.font = value_font
                c.alignment = center
                row += 1
            row += 1

        # عنوان + المورد
        if sections.get("info", True):
            inv_no = invoice.get("invoice_no", "")
            inv_date = invoice.get("invoice_date", "")

            ws.merge_cells(start_row=row, start_column=1,
                           end_row=row, end_column=num_cols)
            t = f"{'فاتورة رقم' if is_ar else 'Invoice No'}: {inv_no}   •   {'التاريخ' if is_ar else 'Date'}: {inv_date}"
            c = ws.cell(row=row, column=1, value=t)
            c.font = section_font
            c.alignment = center
            row += 2

            field_values = {
                "name":       invoice.get("supplier_name", "—"),
                "code":       invoice.get("supplier_code", "—"),
                "address":    invoice.get("supplier_address", "—"),
                "phone":      invoice.get("supplier_phone", "—"),
                "email":      invoice.get("supplier_email", "—"),
                "invoice_no": invoice.get("supplier_invoice_no") or "—",
            }
            field_labels = {
                "name":       "الاسم"              if is_ar else "Name",
                "code":       "الكود"              if is_ar else "Code",
                "address":    "العنوان"            if is_ar else "Address",
                "phone":      "الهاتف"             if is_ar else "Phone",
                "email":      "البريد"             if is_ar else "Email",
                "invoice_no": "رقم فاتورة المورد" if is_ar else "Supplier Inv No",
            }

            if supplier_fields:
                for key in supplier_fields:
                    label = field_labels.get(key, key)
                    val = field_values.get(key, "—")

                    lc = ws.cell(row=row, column=1, value=label)
                    lc.font = label_font
                    lc.alignment = right if is_ar else left
                    lc.fill = info_fill
                    lc.border = border

                    ws.merge_cells(start_row=row, start_column=2,
                                   end_row=row, end_column=num_cols)
                    vc = ws.cell(row=row, column=2, value=str(val))
                    vc.font = value_font
                    vc.alignment = right if is_ar else left
                    vc.fill = info_fill
                    vc.border = border

                    for col in range(2, num_cols + 1):
                        ws.cell(row=row, column=col).border = border
                    row += 1
                row += 1

        # جدول الأصناف
        if sections.get("items", True):
            ws.merge_cells(start_row=row, start_column=1,
                           end_row=row, end_column=num_cols)
            c = ws.cell(row=row, column=1,
                        value=("الأصناف" if is_ar else "Products"))
            c.font = section_font
            c.alignment = center
            row += 1

            col_map = {c[0]: c for c in INVOICE_ITEM_COLUMNS}
            for idx, key in enumerate(columns):
                if key in col_map:
                    lbl = col_map[key][2] if is_ar else col_map[key][1]
                    hc = ws.cell(row=row, column=idx + 1, value=lbl)
                    hc.font = header_font
                    hc.fill = primary_fill
                    hc.alignment = center
                    hc.border = border
            if "layout" in columns and "image" in columns:
                first_col, last_col = sorted((columns.index("layout") + 1, columns.index("image") + 1))
                ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
                ws.cell(row, first_col).value = "layout"
            row += 1

            for it in invoice.get("items", []):
                for idx, key in enumerate(columns):
                    if key == "image":
                        cell = ws.cell(row=row, column=idx + 1)
                        cell.alignment = center
                        cell.border = border
                        image_path = self._get_product_image_path(it)
                        if image_path:
                            picture = ExcelImage(str(image_path))
                            scale = min(48 / picture.width, 40 / picture.height)
                            picture.width *= scale
                            picture.height *= scale
                            picture.anchor = cell.coordinate
                            ws.add_image(picture)
                            ws.row_dimensions[row].height = 32
                        else:
                            cell.value = "—"
                        continue
                    val = self._get_item_column_value(it, key)
                    vc = ws.cell(row=row, column=idx + 1, value=val)
                    vc.font = value_font
                    vc.alignment = center
                    vc.border = border
                row += 1
            quantity_total = sum(float(item.get("quantity") or 0) for item in invoice.get("items", []))
            dollar_total = sum(float(item.get("total_usd") or 0) for item in invoice.get("items", []))
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
            total_label = "Total Qty" if is_ar else "Total Quantity"
            dollar_label = "إجمالي القيمة بالدولار" if is_ar else "Total Value (USD)"
            cell = ws.cell(row, 1, f"{total_label}: {quantity_total:g}    |    {dollar_label}: $ {dollar_total:,.2f}")
            cell.font = kpi_font
            cell.fill = section_fill
            cell.alignment = right if is_ar else left
            row += 2




        export_notes = "\n\n".join(str(invoice.get(key) or "").strip() for key in ("notes", "notes2") if str(invoice.get(key) or "").strip())
        if sections.get("footer", True) and export_notes:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
            cell = ws.cell(row, 1, export_notes)
            cell.font = Font(size=7, name="Cairo")
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            ws.row_dimensions[row].height = max(18, 10 * len(export_notes.splitlines()))
            row += 2

        split_column = max(2, num_cols // 2)
        def write_list_row(label, amount, bold=False, number_format="#,##0.00"):
            nonlocal row
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=split_column)
            ws.merge_cells(start_row=row, start_column=split_column + 1, end_row=row, end_column=num_cols)
            lc = ws.cell(row, 1, label)
            lc.font = kpi_font if bold else value_font
            lc.alignment = right if is_ar else left
            vc = ws.cell(row, split_column + 1, float(amount or 0))
            vc.font = kpi_font if bold else value_font
            vc.alignment = left if is_ar else right
            vc.number_format = number_format
            for column in range(1, num_cols + 1):
                ws.cell(row, column).border = Border(bottom=thin)
            ws.row_dimensions[row].height = 24
            row += 1

        def write_list_title(title):
            nonlocal row
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
            cell = ws.cell(row, 1, title)
            cell.font = section_font
            cell.alignment = right if is_ar else left
            ws.row_dimensions[row].height = 28
            row += 1

        for key, title in (
            ("payments", "المدفوعات" if is_ar else "Payments"),
            ("discounts", "الخصومات" if is_ar else "Discounts"),
            ("expenses", "المصروفات" if is_ar else "Expenses"),
        ):
            if sections.get(key, True):
                write_list_title(title)
                amounts = invoice.get(key) or []
                for item in amounts:
                    write_list_row(item.get("label") or "—", item.get("amount", 0))
                write_list_row("الإجمالي" if is_ar else "Total", sum(float(item.get("amount") or 0) for item in amounts), True)
                row += 1

        if sections.get("summary", True):
            write_list_title("الملخص والأرباح" if is_ar else "Summary and Profit")
            summary_rows = [
                (("إجمالي الدفعات" if is_ar else "Total Payments"),
                 invoice.get("total_payments", 0)),
                (("إجمالي الخصومات" if is_ar else "Total Discounts"),
                 -invoice.get("total_discounts", 0)),
                (("صافي المدفوعات" if is_ar else "Net Payments"),
                 invoice.get("net_payments", 0)),
                (("إجمالي المصاريف" if is_ar else "Total Expenses"),
                 invoice.get("total_expenses", 0)),
                (("صافي التكلفة" if is_ar else "Net Cost (LYD)"),
                 invoice.get("net_cost", 0)),
                (("القيمة بالدولار" if is_ar else "Total Value (USD)"),
                 invoice.get("total_usd", 0)),
                (("سعر الصرف" if is_ar else "Exchange Rate"),
                 invoice.get("exchange_rate", 0)),
                (("إجمالي البيع" if is_ar else "Total Sales (LYD)"),
                 invoice.get("total_sales", 0)),
                (("إجمالي الربح" if is_ar else "Total Profit (LYD)"),
                 invoice.get("total_profit", 0)),
                (("هامش الربح %" if is_ar else "Profit Margin %"),
                 invoice.get("profit_margin", 0)),
            ]


            for index, (label, val) in enumerate(summary_rows):
                number_format = "0.0000" if index == 6 else '0.00" %"' if index == 9 else '#,##0.00" USD"' if index == 5 else '#,##0.00" LYD"'
                write_list_row(label, val, index in (4, 7, 8), number_format)
            row += 1

        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "landscape" if num_cols > 8 else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_area = f"A1:{get_column_letter(num_cols)}{row}"
        ws.column_dimensions["A"].width = 24
        for col_idx in range(2, num_cols + 1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = 18

        _save_literal_workbook(wb, file_path)
        return file_path

    # --------------------------------------------------------
    # Helper
    # --------------------------------------------------------
    @staticmethod
    def _order_pdf_invoice_columns(columns: list[str]) -> list[str]:
        """Keep PDF columns stable, placing image between layout and pallet."""
        requested = set(columns)
        order = [
            "product_name", "pattern", "category", "ah", "l", "w", "h", "oh",
            "layout", "image", "pallet", "quantity", "purchase_usd",
            "purchase_dzd", "sale_dzd", "total_usd", "total_dzd",
            "total_sale", "unit_profit", "total_profit",
        ]
        return [key for key in order if key in requested]

    def _get_product_image_path(self, item: dict) -> Path | None:
        filename = item.get("product_image") or item.get("layout_image")
        if not filename:
            return None
        from app.services.product_service import product_service
        return product_service.get_image_path(filename)

    def _get_item_column_value(self, item: dict, key: str, pdf=False) -> str:
        if pdf and key in ("ah", "l", "w", "h", "oh", "layout", "pallet", "quantity",
                           "purchase_usd", "purchase_dzd", "sale_dzd", "total_usd", "total_dzd",
                           "total_sale", "unit_profit", "total_profit"):
            value = item.get("product_" + key) if key in ("ah", "l", "w", "h", "oh", "layout") else item.get(key)
            if key == "pallet" and value is None:
                value = item.get("product_pallet")
            return _pdf_number(value, 2 if key in ("purchase_usd", "total_usd") else 0)
        if key == "pattern":
            return item.get("product_pattern") or "—"
        if key == "product_name":
            return item.get("product_code") or item.get("product_name", "—")
        if key == "code":
            return item.get("product_code", "—")
        if key == "category":
            return item.get("product_category", "—")
        if key == "ah":
            return self._fmt_num(item.get("product_ah"))
        if key == "l":
            return self._fmt_num(item.get("product_l"))
        if key == "w":
            return self._fmt_num(item.get("product_w"))
        if key == "h":
            return self._fmt_num(item.get("product_h"))
        if key == "oh":
            return self._fmt_num(item.get("product_oh"))
        if key == "layout":
            return self._fmt_num(item.get("product_layout"))
        if key == "pallet":
            pallet = item.get("pallet")
            if pallet is None:
                pallet = item.get("product_pallet")
            return self._fmt_num(pallet)
        if key == "quantity":
            return self._fmt_num(item.get("quantity"))
        if key in ("purchase_usd", "purchase_dzd", "sale_dzd",
                   "total_usd", "total_dzd", "total_sale",
                   "unit_profit", "total_profit"):
            try:
                return f"{float(item.get(key, 0)):,.2f}"
            except (TypeError, ValueError):
                return "0.00"
        return str(item.get(key, "—"))

    def _fmt_num(self, v) -> str:
        if v is None or v == "":
            return "—"
        try:
            f = float(v)
            return str(int(f)) if f.is_integer() else f"{f:g}"
        except (TypeError, ValueError):
            return str(v)

    def _is_arabic(self) -> bool:
        try:
            return (settings.get("language") or "en") == "ar"
        except Exception:
            return False


export_service = ExportService()
