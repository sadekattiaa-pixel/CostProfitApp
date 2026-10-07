import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.db import db
from app.services.export_service import export_service
from app.services.product_service import ProductError, product_service
from app.services.settings_service import settings
from app.services.supplier_service import SupplierError, supplier_service


class SupplierProductServiceTests(unittest.TestCase):
    def setUp(self):
        self.db_path = Path(__file__).with_name("_supplier_product_test.db")
        self.export_path = Path(__file__).with_name("_products_export_test.xlsx")
        db.connect(self.db_path)
        settings.clear_cache()

    def tearDown(self):
        db.close()
        settings.clear_cache()
        for path in (
            self.db_path,
            Path(str(self.db_path) + "-wal"),
            Path(str(self.db_path) + "-shm"),
            self.export_path,
        ):
            if path.exists():
                path.unlink()

    def test_supplier_crud_and_duplicate_code_validation(self):
        supplier_id = supplier_service.create("S-001", "Supplier One", phone="0910000000")
        supplier_service.update(supplier_id, "S-001", "Updated Supplier", phone="0920000000")

        supplier = supplier_service.get_by_id(supplier_id)
        self.assertEqual(supplier["name"], "Updated Supplier")
        self.assertEqual(supplier["phone"], "0920000000")
        with self.assertRaises(SupplierError):
            supplier_service.create("S-001", "Duplicate")

        supplier_service.delete(supplier_id)
        self.assertIsNone(supplier_service.get_by_id(supplier_id))

    def test_product_code_normalization_and_supplier_delete_protection(self):
        supplier_id = supplier_service.create("S-001", "Supplier")
        product_id = product_service.create(
            "P-0001", "Premium Battery", "Batteries", supplier_id, pallet=20,
        )

        with self.assertRaises(ProductError):
            product_service.create(
                "  p-0001  ", "Unused name", "Batteries", supplier_id,
            )
        with self.assertRaises(SupplierError):
            supplier_service.delete(supplier_id)

        product_service.delete(product_id)
        supplier_service.delete(supplier_id)
        self.assertEqual(product_service.count(), 0)
        self.assertEqual(supplier_service.count(), 0)

    def test_products_excel_export_contains_expected_values(self):
        supplier_id = supplier_service.create("S-001", "Supplier")
        product_service.create(
            "P-0001", "Premium Battery", "Batteries", supplier_id, ah=70, pallet=20,
        )
        products = product_service.get_all()

        result = export_service.export_products_excel(products, str(self.export_path), "en")
        self.assertEqual(Path(result), self.export_path)

        workbook = load_workbook(self.export_path, read_only=True, data_only=True)
        try:
            sheet = workbook["Products"]
            self.assertEqual(sheet["A2"].value, "P-0001")
            self.assertEqual(sheet["B2"].value, "Batteries")
            self.assertEqual(sheet["C2"].value, "Supplier")
            self.assertEqual(sheet.max_column, 11)
            self.assertNotIn("Product Name", [cell.value for cell in sheet[1]])
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
